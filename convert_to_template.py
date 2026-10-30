import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
import datetime

def format_date_value(val):
    if val is None:
        return ''
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime('%d/%m/%Y')
    return str(val).strip()

def process_excel(input_file, output_file):
    wb_in = openpyxl.load_workbook(input_file, data_only=True)
    ws_in = wb_in.active

    # Data structures
    customers_dict = {}  # k: customer_code (Col B/1), v: customer_name (Col C/2)
    debts = [] # list of dicts

    current_customer_code = None
    
    # Iterate through rows starting from data rows
    # The header seems to be around row 10-11
    # We will look for rows where Col 1 (MÃ KHÁCH HÀNG) or Col 2 (TÊN KHÁCH HÀNG) exists.
    for i, row in enumerate(ws_in.iter_rows(values_only=True)):
        if i < 10: # Skip early headers
            continue
            
        stt = row[0]
        ma_kh = row[1]
        ten_kh = row[2]
        
        if ma_kh and ten_kh and str(stt) not in ('0', 'None', ''):
            customers_dict[ma_kh] = ten_kh
            current_customer_code = ma_kh
            
        # Try to infer current customer if not explicitly set but ma_kh is present
        if ma_kh and not ten_kh:
            current_customer_code = ma_kh
            
        # Is this an invoice row?
        noi_dung = row[7] # Col H
        hd_so = row[8] # Col I
        hd_ngay = row[9] # Col J
        so_tien_phai_thu = row[10] # Col K
        thoi_han_tt = row[11] # Col L
        han_thu = row[12] # Col M
        da_thu_tien = row[13] # Col N
        con_lai_phai_thu = row[14] # Col O
        qua_han_tt = row[15] # Col P
        
        # Check if it's an overdue invoice
        # It must have hd_so, and qua_han_tt must be > 0 (or at least valid)
        if hd_so:
            try:
                days_overdue = int(qua_han_tt) if qua_han_tt else 0
            except ValueError:
                days_overdue = 0
                
            if days_overdue > 0:
                cust_name = customers_dict.get(current_customer_code, "")
                debts.append({
                    'customer_name': cust_name,
                    'customer_code': current_customer_code,
                    'noi_dung': noi_dung,
                    'hd_so': hd_so,
                    'hd_ngay': hd_ngay,
                    'so_tien_phai_thu': so_tien_phai_thu,
                    'thoi_han_tt': thoi_han_tt,
                    'han_thu': han_thu,
                    'da_thu_tien': da_thu_tien,
                    'con_lai_phai_thu': con_lai_phai_thu,
                    'qua_han_tt': qua_han_tt
                })

    wb_in.close()

    # Now create the template workbook
    wb_out = openpyxl.Workbook()

    # Sheet 1: HuongDan
    ws0 = wb_out.active
    ws0.title = 'HuongDan'
    ws0.sheet_properties.tabColor = '4472C4'
    instructions = [
        ['HƯỚNG DẪN SỬ DỤNG TEMPLATE CÔNG NỢ'],
        ['1. Sheet "DanhSach": Danh sách khách hàng và thông tin email gởi.'],
        ['   Lưu ý: Rê chuột vào tiêu đề cột (dòng 1) để xem chú thích chi tiết cho từng trường.'],
        ['2. Sheet "CongNo": Chi tiết từng khoản nợ (được tự động trích xuất từ dữ liệu gốc).'],
    ]
    for row_data in instructions:
        ws0.append(row_data)

    # Styles
    hf = Font(bold=True, color='FFFFFF', size=11)
    hfill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    ha = Alignment(horizontal='center', vertical='center', wrap_text=True)
    tb = Border(left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))
    ef = Font(italic=True, color='808080', size=10)
    efill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

    # Sheet 2: DanhSach
    ws1 = wb_out.create_sheet('DanhSach')
    ws1.sheet_properties.tabColor = '00B050'
    headers1 = [
        ('customer_name', 'Tên khách hàng\n(Sẽ được đưa vào sau chữ Kính gửi: trong email)', 35),
        ('customer_code', 'Mã khách hàng\n(Dùng để tự động tìm file đính kèm chứa mã này)', 20),
        ('short_name', 'Tên viết tắt của khách hàng\n(Sẽ dùng trên tiêu đề của file/email)', 25),
        ('email_to', 'Email người nhận chính\n(Ngăn cách nhiều email bằng dấu chấm phẩy ;)', 30),
        ('email_cc', 'Email người nhận CC\n(Ngăn cách nhiều email bằng dấu chấm phẩy ;)', 30),
        ('email_bcc', 'Email BCC (ẩn danh)\n(Ngăn cách nhiều email bằng dấu chấm phẩy ;)', 30),
        ('debt_date', 'Ngày chốt số công nợ\nĐịnh dạng: dd/mm/yyyy', 20),
        ('payment_deadline', 'Hạn yêu cầu thanh toán / nhắc nợ\nĐịnh dạng: dd/mm/yyyy', 25),
        ('company_name', 'Tên công ty / Chi nhánh\n(Sẽ hiển thị trong bảng chi tiết email)', 30),
        ('attachment_dir', 'Đường dẫn thư mục chứa file riêng của KH này\n(Ví dụ: D:\\HoaDon)', 30),
        ('attachment_file', 'Tên file riêng của KH này\n(Ví dụ: HD001.pdf)', 25),
    ]
    for col, (key, desc, w) in enumerate(headers1, 1):
        c = ws1.cell(row=1, column=col, value=key)
        c.font = hf; c.fill = hfill; c.alignment = ha; c.border = tb
        ws1.column_dimensions[get_column_letter(col)].width = w
        c.comment = Comment(desc, 'Hướng dẫn')

    # Find unique customer names and codes that have debts
    unique_customers = {} # name -> code
    for d in debts:
        if d['customer_name']:
            unique_customers[d['customer_name']] = d.get('customer_code', '')

    today_str = datetime.datetime.now().strftime('%d/%m/%Y')
    
    for ri, (cust_name, cust_code) in enumerate(unique_customers.items(), 2):
        ws1.cell(row=ri, column=1, value=cust_name)
        ws1.cell(row=ri, column=2, value=cust_code)
        ws1.cell(row=ri, column=3, value=cust_code) # default short_name to code
        ws1.cell(row=ri, column=7, value=today_str) # debt_date
        ws1.cell(row=ri, column=8, value='') # payment_deadline blank

    # Sheet 3: CongNo
    ws2 = wb_out.create_sheet('CongNo')
    ws2.sheet_properties.tabColor = 'FFC000'
    headers2 = [
        ('customer_name', 'Tên Khách Hàng\n(Phải khớp với tên bên sheet DanhSach)', 30),
        ('noi_dung', 'Nội dung / Sản phẩm của đơn hàng\n[Tương ứng Cột H của file gốc]', 25),
        ('hd_so', 'Số hóa đơn gốc\n[Tương ứng Cột I]', 15),
        ('hd_ngay', 'Ngày xuất hóa đơn\n[Tương ứng Cột J]', 18),
        ('so_tien_phai_thu', 'Số tiền phải thu (VNĐ)\n[Tương ứng Cột K]', 22),
        ('thoi_han_tt', 'Thời hạn thanh toán đã thoả thuận (Ngày)\n[Tương ứng Cột L]', 20),
        ('han_thu', 'Ngày đến hạn thu\n[Tương ứng Cột M]', 18),
        ('da_thu_tien', 'Số tiền đã thu được\n[Tương ứng Cột N]', 20),
        ('con_lai_phai_thu', 'Khoản nợ còn lại phải thu\n[Tương ứng Cột O]', 22),
        ('qua_han_tt', 'Số ngày đã quá hạn thanh toán\n[Tương ứng Cột P]', 20),
    ]
    for col, (key, desc, w) in enumerate(headers2, 1):
        c = ws2.cell(row=1, column=col, value=key)
        c.font = hf; c.fill = hfill; c.alignment = ha; c.border = tb
        ws2.column_dimensions[get_column_letter(col)].width = w
        c.comment = Comment(desc, 'Hướng dẫn')

    for ri, d in enumerate(debts, 2):
        exs = [
            d['customer_name'],
            d['noi_dung'],
            d['hd_so'],
            format_date_value(d['hd_ngay']),
            d['so_tien_phai_thu'],
            d['thoi_han_tt'],
            format_date_value(d['han_thu']),
            d['da_thu_tien'],
            d['con_lai_phai_thu'],
            d['qua_han_tt']
        ]
        for col, val in enumerate(exs, 1):
            c = ws2.cell(row=ri, column=col, value=val)
            c.border = tb

    wb_out.save(output_file)
    wb_out.close()
    print(f"Successfully converted to {output_file} with {len(debts)} debts over {len(unique_customers)} customers.")

if __name__ == '__main__':
    import sys
    import shutil
    import tempfile
    import os
    
    input_f = sys.argv[1] if len(sys.argv) > 1 else 'Mau_Bravo8.xlsx'
    output_f = sys.argv[2] if len(sys.argv) > 2 else 'Template_SenEmail.xlsx'
    
    if os.path.exists(input_f):
        temp_file = os.path.join(tempfile.gettempdir(), 'temp_convert_input.xlsx')
        shutil.copy(input_f, temp_file)
        try:
            process_excel(temp_file, output_f)
        finally:
            if os.path.exists(temp_file):
                try:
                    os.remove(temp_file)
                except:
                    pass
    else:
        print(f"File not found: {input_f}. Usage: python convert_to_template.py <input_excel> <output_excel>")
