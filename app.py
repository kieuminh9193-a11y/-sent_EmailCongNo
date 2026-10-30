# -*- coding: utf-8 -*-
"""
Flask Backend - Debt Collection Email Tool with Mail Merge.
"""
import os
import sys
import json
import datetime
from flask import Flask, render_template, request, jsonify, send_file
from openpyxl import Workbook, load_workbook
from email_template import merge_template, DEFAULT_TEMPLATE
from outlook_sender import create_outlook_email, check_outlook_available
from convert_to_template import process_excel
from license_check import check_license, get_license_info
# Detect PyInstaller frozen mode
if getattr(sys, 'frozen', False):
    # Running as bundled exe: templates/static are in temp _MEIPASS folder
    BUNDLE_DIR = sys._MEIPASS
    # User data (uploads, settings, logs) lives next to the exe
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BUNDLE_DIR = os.path.dirname(os.path.abspath(__file__))
    BASE_DIR = BUNDLE_DIR

app = Flask(__name__,
            template_folder=os.path.join(BUNDLE_DIR, 'templates'),
            static_folder=os.path.join(BUNDLE_DIR, 'static'))
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

APP_VERSION = '1.1.4'
APP_AUTHOR = 'DebtMail Pro'

UPLOAD_DIR = os.path.join(BASE_DIR, 'uploads')
SETTINGS_FILE = os.path.join(BASE_DIR, 'settings.json')
LOG_FILE = os.path.join(BASE_DIR, 'send_log.json')

os.makedirs(UPLOAD_DIR, exist_ok=True)


# ============================================================
# SEND LOG
# ============================================================
def load_send_log():
    if os.path.exists(LOG_FILE):
        with open(LOG_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}


def save_send_log(log):
    with open(LOG_FILE, 'w', encoding='utf-8') as f:
        json.dump(log, f, ensure_ascii=False, indent=2)


def record_send(customer_name, email_to, subject):
    """Record a send event for a customer."""
    log = load_send_log()
    if customer_name not in log:
        log[customer_name] = {'count': 0, 'history': []}
    log[customer_name]['count'] += 1
    log[customer_name]['history'].append({
        'date': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'to': email_to,
        'subject': subject
    })
    # Keep only last 50 history entries per customer
    log[customer_name]['history'] = log[customer_name]['history'][-50:]
    save_send_log(log)
    return log[customer_name]['count']

# ============================================================
# SETTINGS (simplified)
# ============================================================
DEFAULT_SETTINGS = {
    'sender_email': '',
    'company_name': '',
    'subject_format': 'DebtCollection_{short_name} || CONG NO QUA HAN PHAI THU DEN  {debt_date_dot}',
    'email_template': DEFAULT_TEMPLATE,
    'signature': '''<p>&nbsp;</p>
<p style="font-family:Arial;color:#16A2DC;">Trân trọng / Best Regards,</p>
<p><b style="font-size:12pt;font-family:'Palatino Linotype';color:#16A2DC;">[Họ và Tên] | [Chức vụ]</b></p>
<p style="font-family:Arial;color:#16A2DC;">Điện thoại: [Số điện thoại]</p>
<p><b style="font-family:Arial;color:#16A2DC;">[TÊN CÔNG TY]</b></p>
<p style="font-family:Arial;font-size:9pt;color:#555;">Địa chỉ: [Địa chỉ công ty / văn phòng]<br>Hotline: [Hotline]</p>''',
    'send_mode': 'display',
    'kh_attachments_dir': '',
    'manual_aliases': '',
}


def load_settings():
    if os.path.exists(SETTINGS_FILE):
        with open(SETTINGS_FILE, 'r', encoding='utf-8') as f:
            saved = json.load(f)
            settings = DEFAULT_SETTINGS.copy()
            settings.update(saved)
            return settings
    return DEFAULT_SETTINGS.copy()


def save_settings(settings):
    with open(SETTINGS_FILE, 'w', encoding='utf-8') as f:
        json.dump(settings, f, ensure_ascii=False, indent=2)


def format_date_value(val):
    if val is None:
        return ''
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime('%d/%m/%Y')
    return str(val).strip()


def format_number_value(val):
    if val is None:
        return ''
    if isinstance(val, (int, float)):
        return val if val != 0 else ''
    s = str(val).strip().replace(',', '')
    try:
        num = float(s)
        return num if num != 0 else ''
    except ValueError:
        return val


def read_excel_data(filepath):
    wb = load_workbook(filepath, data_only=True)
    
    customers = []
    if 'DanhSach' in wb.sheetnames:
        ws = wb['DanhSach']
        headers = [cell.value for cell in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            customer = {}
            for i, header in enumerate(headers):
                if header and i < len(row):
                    key = header.strip()
                    val = row[i]
                    if key in ('debt_date', 'payment_deadline'):
                        val = format_date_value(val)
                    customer[key] = val if val is not None else ''
            customers.append(customer)
    
    debts = {}
    if 'CongNo' in wb.sheetnames:
        ws = wb['CongNo']
        headers = [cell.value for cell in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            item = {}
            for i, header in enumerate(headers):
                if header and i < len(row):
                    key = header.strip()
                    val = row[i]
                    if key in ('hd_ngay', 'han_thu'):
                        val = format_date_value(val)
                    elif key in ('so_tien_phai_thu', 'da_thu_tien', 'con_lai_phai_thu'):
                        val = format_number_value(val)
                    elif key in ('thoi_han_tt', 'qua_han_tt'):
                        if isinstance(val, float) and val == int(val):
                            val = int(val)
                        val = val if val else ''
                    else:
                        val = val if val is not None else ''
                    item[key] = val
            
            cust_name = str(row[0]).strip()
            if cust_name not in debts:
                debts[cust_name] = []
            debts[cust_name].append(item)
    
    wb.close()
    return {'customers': customers, 'debts': debts}


def create_sample_excel(filepath):
    """Create Excel template for user to fill in."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment

    wb = Workbook()

    # Sheet 1: HuongDan
    ws0 = wb.active
    ws0.title = 'HuongDan'
    ws0.sheet_properties.tabColor = '4472C4'
    instructions = [
        ['HUONG DAN SU DUNG TEMPLATE CONG NO'],
        [''],
        ['1. Sheet "DanhSach": Danh sach khach hang'],
        ['   - customer_name: Ten KH (hien thi sau "Kinh gui:")'],
        ['   - customer_code: Ma KH (dung de tu dong tim va ghep file dinh kem)'],
        ['   - short_name: Ten viet tat (dung trong tieu de email)'],
        ['   - email_to / email_cc / email_bcc: Email nguoi nhan, CC va BCC (ngan bang ;)'],
        ['   - debt_date: Ngay tinh cong no (dd/mm/yyyy)'],
        ['   - payment_deadline: Han thanh toan yeu cau (dd/mm/yyyy)'],
        ['   - company_name: Ten cong ty hien thi trong bang va noi dung'],
        [''],
        ['2. Sheet "CongNo": Chi tiet tung khoan no'],
        ['   - customer_name: PHAI KHOP voi sheet DanhSach'],
        ['   - Cac cot so tien nhap dang so (vd: 241380000)'],
        [''],
        ['3. Trong email template, dung {ten_truong} de chen du lieu:'],
        ['   {customer_name}  - Ten khach hang'],
        ['   {customer_code}  - Ma khach hang'],
        ['   {company_name}   - Ten cong ty'],
        ['   {debt_date}      - Ngay tinh cong no'],
        ['   {payment_deadline} - Han thanh toan'],
        ['   {debt_table}     - Bang cong no (tu dong tao)'],
        ['   {signature}      - Chu ky email'],
        ['   {short_name}     - Ten viet tat'],
        [''],
        ['LUU Y: Dong mau xam la VI DU, hay xoa va nhap data thuc'],
    ]
    for row_data in instructions:
        ws0.append(row_data)
    ws0['A1'].font = Font(bold=True, size=14, color='2F5496')
    for r in range(3, 28):
        ws0.cell(row=r, column=1).font = Font(size=11)
    ws0.column_dimensions['A'].width = 80

    # Styles
    hf = Font(bold=True, color='FFFFFF', size=11)
    hfill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    ha = Alignment(horizontal='center', vertical='center', wrap_text=True)
    tb = Border(left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))
    ef = Font(italic=True, color='808080', size=10)
    efill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

    # Sheet 2: DanhSach
    ws1 = wb.create_sheet('DanhSach')
    ws1.sheet_properties.tabColor = '00B050'
    headers1 = [
        ('customer_name', 'Ten KH (Kinh gui:)', 30),
        ('customer_code', 'Ma KH (de khop file dinh kem)', 15),
        ('short_name', 'Ten viet tat', 15),
        ('email_to', 'Email nguoi nhan', 30),
        ('email_cc', 'Email CC', 30),
        ('email_bcc', 'Email BCC (an danh)', 30),
        ('debt_date', 'Ngay tinh cong no\n(dd/mm/yyyy)', 18),
        ('payment_deadline', 'Han thanh toan\n(dd/mm/yyyy)', 18),
        ('company_name', 'Ten cong ty\n(hien thi trong bang)', 30),
        ('attachment_dir', 'Duong dan thu muc file rieng KH (VD: D:\\HoaDon)', 30),
        ('attachment_file', 'Ten file dinh kem rieng (VD: HD001.pdf)', 25),
    ]
    for col, (key, desc, w) in enumerate(headers1, 1):
        c = ws1.cell(row=1, column=col, value=key)
        c.font = hf; c.fill = hfill; c.alignment = ha; c.border = tb
        ws1.column_dimensions[get_column_letter(col)].width = w
        c.comment = Comment(desc, 'DebtMail Pro')
    ex1 = ['CONG TY CP ABC', 'KH001', 'ABC', 'ketoan@abc.com', 'giamdoc@abc.com',
           'quanly@abc.com', '01/01/2026', '15/01/2026', 'Cong ty Mau ABC', 'D:\\HoaDon', 'HD001.pdf']
    for col, val in enumerate(ex1, 1):
        c = ws1.cell(row=2, column=col, value=val)
        c.font = ef; c.fill = efill; c.border = tb
    ws1.row_dimensions[1].height = 40
    ws1.freeze_panes = 'A2'

    # Sheet 3: CongNo
    ws2 = wb.create_sheet('CongNo')
    ws2.sheet_properties.tabColor = 'FFC000'
    headers2 = [
        ('customer_name', 'Ten KH (khop DanhSach)', 30),
        ('noi_dung', 'Noi dung / San pham', 25),
        ('hd_so', 'So hoa don', 12),
        ('hd_ngay', 'Ngay HD (dd/mm/yyyy)', 15),
        ('so_tien_phai_thu', 'So tien phai thu', 20),
        ('thoi_han_tt', 'Thoi han TT (ngay)', 14),
        ('han_thu', 'Han thu (dd/mm/yyyy)', 15),
        ('da_thu_tien', 'Da thu tien', 18),
        ('con_lai_phai_thu', 'Con lai phai thu', 18),
        ('qua_han_tt', 'Qua han TT (ngay)', 14),
    ]
    for col, (key, desc, w) in enumerate(headers2, 1):
        c = ws2.cell(row=1, column=col, value=key)
        c.font = hf; c.fill = hfill; c.alignment = ha; c.border = tb
        ws2.column_dimensions[get_column_letter(col)].width = w
        c.comment = Comment(desc, 'DebtMail Pro')
    exs = [
        ['CONG TY CP ABC', 'LABRAFAC MC60', '1718', '30/12/2025', 241380000, 30, '29/01/2026', 100000000, 141380000, 68],
        ['CONG TY CP ABC', 'STEARIC ACID BP', '203', '10/02/2026', 6750000, 30, '12/03/2026', 0, 6750000, 26],
    ]
    for ri, rd in enumerate(exs, 2):
        for col, val in enumerate(rd, 1):
            c = ws2.cell(row=ri, column=col, value=val)
            c.font = ef; c.fill = efill; c.border = tb
    ws2.row_dimensions[1].height = 40
    ws2.freeze_panes = 'A2'

    wb.save(filepath)
    wb.close()


# ============================================================
# ROUTES
# ============================================================

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/settings', methods=['GET'])
def get_settings():
    return jsonify(load_settings())

@app.route('/api/settings', methods=['POST'])
def update_settings():
    data = request.json
    settings = load_settings()
    settings.update(data)
    save_settings(settings)
    return jsonify({'success': True, 'settings': settings})

@app.route('/api/outlook-check', methods=['GET'])
def outlook_check():
    return jsonify(check_outlook_available())

@app.route('/api/upload', methods=['POST'])
def upload_excel():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'No file found'}), 400
    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'success': False, 'message': 'Only .xlsx files accepted'}), 400
    filepath = os.path.join(UPLOAD_DIR, 'cong_no_data.xlsx')
    file.save(filepath)
    try:
        data = read_excel_data(filepath)
        return jsonify({
            'success': True,
            'message': f'Upload OK. Found {len(data["customers"])} customers.',
            'customers': data['customers']
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'Error: {str(e)}'}), 400

@app.route('/api/customers', methods=['GET'])
def get_customers():
    filepath = os.path.join(UPLOAD_DIR, 'cong_no_data.xlsx')
    if not os.path.exists(filepath):
        return jsonify({'customers': [], 'message': 'No Excel file uploaded'})
    try:
        data = read_excel_data(filepath)
        for cust in data['customers']:
            name = cust.get('customer_name', '')
            items = data['debts'].get(name, [])
            cust['debt_count'] = len(items)
            cust['total_debt'] = sum(
                float(str(item.get('con_lai_phai_thu', 0) or 0).replace(',', ''))
                for item in items
            )
        return jsonify({'customers': data['customers']})
    except Exception as e:
        return jsonify({'customers': [], 'message': str(e)})

@app.route('/api/upload-b8', methods=['POST'])
def upload_b8():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'No file found'}), 400
    file = request.files['file']
    output_mode = request.form.get('output_mode', 'import') # 'import' or 'download'

    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'success': False, 'message': 'Only .xlsx files accepted'}), 400
        
    temp_in = os.path.join(UPLOAD_DIR, 'temp_b8_input.xlsx')
    output_path = os.path.join(UPLOAD_DIR, 'cong_no_data.xlsx')
    
    # If downloading, we shouldn't overwrite the main data file
    if output_mode == 'download':
        output_path = os.path.join(UPLOAD_DIR, 'converted_template.xlsx')
    
    try:
        file.save(temp_in)
        # Convert B8 to Template
        process_excel(temp_in, output_path)
        
        # Cleanup temp
        if os.path.exists(temp_in):
            os.remove(temp_in)
            
        if output_mode == 'download':
            return send_file(output_path, as_attachment=True, download_name='converted_template.xlsx')

        data = read_excel_data(output_path)
        return jsonify({
            'success': True,
            'message': f'Chuyển đổi thành công! Tìm thấy {len(data["customers"])} khách hàng.',
            'customers': data['customers']
        })
    except Exception as e:
        if os.path.exists(temp_in):
            os.remove(temp_in)
        return jsonify({'success': False, 'message': f'Lỗi chuyển đổi: {str(e)}'}), 400

@app.route('/api/clear-data', methods=['POST'])
def clear_data():
    filepath = os.path.join(UPLOAD_DIR, 'cong_no_data.xlsx')
    if os.path.exists(filepath):
        try:
            os.remove(filepath)
            return jsonify({'success': True, 'message': 'Đã xóa danh sách.'})
        except Exception as e:
            return jsonify({'success': False, 'message': str(e)}), 400
    return jsonify({'success': True, 'message': 'Danh sách trống.'})

@app.route('/api/customer/<path:name>', methods=['GET'])
def get_customer_detail(name):
    filepath = os.path.join(UPLOAD_DIR, 'cong_no_data.xlsx')
    if not os.path.exists(filepath):
        return jsonify({'success': False, 'message': 'No file'})
    data = read_excel_data(filepath)
    customer = None
    for c in data['customers']:
        if c.get('customer_name', '') == name:
            customer = c
            break
    if not customer:
        return jsonify({'success': False, 'message': f'Not found: {name}'})
        
    # Auto-match attachments from directory based on customer code
    matched_attachments = []
    settings = load_settings()
    attachments_dir = settings.get('kh_attachments_dir', '')
    cust_code = str(customer.get('customer_code', '')).strip()
    
    if attachments_dir and os.path.exists(attachments_dir) and cust_code:
        try:
            for filename in os.listdir(attachments_dir):
                file_path = os.path.join(attachments_dir, filename)
                if os.path.isfile(file_path):
                    # Case-insensitive match on customer code
                    if cust_code.lower() in filename.lower():
                        matched_attachments.append({
                            'name': filename,
                            'path': file_path
                        })
        except Exception as e:
            print(f"Error scanning attachments directory: {e}")

    # Match Excel-specified individual attachments
    excel_dir = str(customer.get('attachment_dir', '')).strip()
    excel_file = str(customer.get('attachment_file', '')).strip()
    if excel_file:
        full_excel_path = ""
        # If absolute path, use it directly
        if os.path.isabs(excel_file):
            full_excel_path = excel_file
        elif excel_dir:
            full_excel_path = os.path.join(excel_dir, excel_file)
            
        if full_excel_path and os.path.exists(full_excel_path) and os.path.isfile(full_excel_path):
            filename = os.path.basename(full_excel_path)
            # Avoid duplicate paths
            if not any(att['path'] == full_excel_path for att in matched_attachments):
                matched_attachments.append({
                    'name': filename,
                    'path': full_excel_path
                })

    return jsonify({
        'success': True,
        'customer': customer,
        'debt_items': data['debts'].get(name, []),
        'matched_attachments': matched_attachments
    })

@app.route('/api/preview', methods=['POST'])
def preview_email():
    data = request.json
    settings = load_settings()
    
    template = data.get('template', settings.get('email_template', DEFAULT_TEMPLATE))
    signature = data.get('signature', settings.get('signature', ''))
    customer = data.get('customer', {})
    debt_items = data.get('debt_items', [])
    company_name = customer.get('company_name', settings.get('company_name', ''))

    # Build merge fields from customer data + settings
    fields = {
        'customer_name': customer.get('customer_name', ''),
        'short_name': customer.get('short_name', ''),
        'company_name': company_name,
        'debt_date': customer.get('debt_date', ''),
        'payment_deadline': customer.get('payment_deadline', ''),
        'signature': signature,
    }

    html = merge_template(template, fields, debt_items, company_name)
    return jsonify({'success': True, 'html': html})

@app.route('/api/send', methods=['POST'])
def send_email():
    data = request.json
    settings = load_settings()
    result = create_outlook_email(
        to=data.get('to', ''),
        cc=data.get('cc', ''),
        subject=data.get('subject', ''),
        html_body=data.get('html_body', ''),
        attachments=data.get('attachments', []),
        send_mode=data.get('send_mode', settings.get('send_mode', 'display')),
        from_account=data.get('from_account', settings.get('sender_email', '')),
        bcc=data.get('bcc', '')
    )
    if result.get('success'):
        count = record_send(
            data.get('customer_name', ''),
            data.get('to', ''),
            data.get('subject', '')
        )
        result['remind_count'] = count
    return jsonify(result)

@app.route('/api/send-batch', methods=['POST'])
def send_batch():
    import time
    data = request.json
    emails = data.get('emails', [])
    settings = load_settings()
    mode = data.get('send_mode', settings.get('send_mode', 'display'))
    results = []
    for i, em in enumerate(emails):
        r = create_outlook_email(
            to=em.get('to', ''), cc=em.get('cc', ''),
            subject=em.get('subject', ''), html_body=em.get('html_body', ''),
            attachments=em.get('attachments', []), send_mode=mode,
            from_account=em.get('from_account', settings.get('sender_email', '')),
            bcc=em.get('bcc', '')
        )
        if r.get('success'):
            record_send(em.get('customer_name', ''), em.get('to', ''), em.get('subject', ''))
        results.append({'customer': em.get('customer_name', ''), **r})
        # Small delay between Display calls to let Outlook process each draft
        if mode == 'display' and i < len(emails) - 1:
            time.sleep(0.5)
    ok = sum(1 for r in results if r['success'])
    return jsonify({'success': True, 'message': f'Processed {ok}/{len(results)} emails', 'results': results})

@app.route('/api/send-log', methods=['GET'])
def get_send_log():
    return jsonify(load_send_log())

@app.route('/api/about', methods=['GET'])
def about():
    return jsonify({
        'version': APP_VERSION,
        'author': APP_AUTHOR,
        'app_name': 'DebtMail Pro',
        'description': 'Công cụ gửi email nhắc công nợ hàng loạt (Mail Merge)',
    })

@app.route('/api/sample-excel', methods=['GET'])
def download_sample():
    filepath = os.path.join(UPLOAD_DIR, 'mau_cong_no.xlsx')
    create_sample_excel(filepath)
    return send_file(filepath, as_attachment=True, download_name='template_cong_no.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/api/generate-subject', methods=['POST'])
def generate_subject():
    data = request.json
    settings = load_settings()
    fmt = data.get('format', settings.get('subject_format', ''))
    short_name = data.get('short_name', '')
    debt_date = data.get('debt_date', '')
    debt_date_dot = debt_date.replace('/', '.')
    subject = fmt.replace('{short_name}', short_name)\
                 .replace('{debt_date}', debt_date)\
                 .replace('{debt_date_dot}', debt_date_dot)
    return jsonify({'subject': subject})
@app.route('/api/license', methods=['GET'])
def license_status():
    info = get_license_info()
    return jsonify(info or {})

if __name__ == '__main__':
    # License check
    is_valid, msg, days = check_license()
    print("=" * 60)
    print("  DEBTMAIL PRO v1.1.4")
    print("=" * 60)
    if not is_valid:
        print("\n" + msg)
        print("\nApp expired. Contact administrator.")
        input("Press Enter to exit...")
        exit(1)
    else:
        print(f"  License: OK - {days} days remaining")
        print(f"  Open: http://localhost:5000")
        print("=" * 60)
    
    import webbrowser
    webbrowser.open('http://localhost:5000')
    app.run(debug=False, port=5000)
